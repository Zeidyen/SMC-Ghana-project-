"""
Minimal Markdown -> Word (.docx) converter for the manuscript. Handles headings,
paragraphs, bold/italic runs, pipe tables, bullet lists and horizontal rules.

    python md2docx.py <in.md> <out.docx>
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(paragraph, text):
    """Split text on **bold** / *italic* and add formatted runs."""
    for part in re.split(r'(\*\*.+?\*\*|\*.+?\*)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = paragraph.add_run(part[1:-1]); r.italic = True
        else:
            paragraph.add_run(part)


def build(md_path, out_path):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    lines = open(md_path).read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]

        # table block
        if ln.strip().startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[-:\s|]+\|\s*$', lines[i + 1]):
            header = [c.strip() for c in ln.strip().strip('|').split('|')]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = 'Light Grid Accent 1'
            for j, h in enumerate(header):
                add_runs(t.rows[0].cells[j].paragraphs[0], h)
                for r in t.rows[0].cells[j].paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for j, c in enumerate(row[:len(header)]):
                    add_runs(cells[j].paragraphs[0], c)
            doc.add_paragraph()
            continue

        s = ln.strip()
        if s == '---':
            doc.add_paragraph()
        elif s.startswith('### '):
            doc.add_heading(s[4:], level=2)
        elif s.startswith('## '):
            doc.add_heading(s[3:], level=1)
        elif s.startswith('# '):
            h = doc.add_heading('', level=0); add_runs(h, s[2:])
        elif s.startswith('- '):
            add_runs(doc.add_paragraph(style='List Bullet'), s[2:])
        elif re.match(r'^\d+\.\s', s):
            add_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s', '', s))
        elif s == '':
            pass
        else:
            add_runs(doc.add_paragraph(), s)
        i += 1

    doc.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
